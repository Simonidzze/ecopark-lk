from datetime import date, datetime, time, timedelta
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from sqlalchemy import func, or_, select

from .db import make_session_factory
from .expense_allocations import expense_receipt_items, money
from .models import Expense, ExpenseAllocation, ExpenseBudgetCategory
from .payment_report import XLSX_CONTENT_TYPE


def parse_date_filter(value):
    value = str(value or "").strip()
    if not value:
        return "", None
    try:
        return value, date.fromisoformat(value)
    except ValueError:
        return "", None


def expense_filter_values(args):
    date_from, date_from_value = parse_date_filter(args.get("date_from"))
    date_to, date_to_value = parse_date_filter(args.get("date_to"))
    return {
        "q": str(args.get("q") or "").strip(),
        "category": str(args.get("category") or "").strip(),
        "category_label": "",
        "date_from": date_from,
        "date_to": date_to,
        "date_from_value": date_from_value,
        "date_to_value": date_to_value,
    }


def allocated_amount_subquery():
    return (
        select(func.coalesce(func.sum(ExpenseAllocation.amount), 0))
        .where(ExpenseAllocation.expense_id == Expense.id)
        .correlate(Expense)
        .scalar_subquery()
    )


def build_expense_statement(filters):
    statement = select(Expense)
    search = filters.get("q", "")
    category = filters.get("category", "")
    date_from = filters.get("date_from_value")
    date_to = filters.get("date_to_value")

    if search:
        like = f"%{search}%"
        budget_category_expenses = (
            select(ExpenseAllocation.expense_id)
            .join(
                ExpenseBudgetCategory,
                ExpenseBudgetCategory.id == ExpenseAllocation.category_id,
            )
            .where(ExpenseBudgetCategory.name.like(like))
        )
        statement = statement.where(
            or_(
                Expense.document.like(like),
                Expense.number.like(like),
                Expense.expense_category.like(like),
                Expense.counterparty.like(like),
                Expense.purpose.like(like),
                Expense.organization.like(like),
                Expense.id.in_(budget_category_expenses),
            )
        )
    if category == "unallocated":
        statement = statement.where(Expense.amount > allocated_amount_subquery())
    elif category:
        statement = statement.where(
            or_(
                Expense.expense_category == category,
                Expense.id.in_(
                    select(ExpenseAllocation.expense_id).where(
                        ExpenseAllocation.category_id == category
                    )
                ),
            )
        )
    if date_from:
        statement = statement.where(Expense.date >= datetime.combine(date_from, time.min))
    if date_to:
        statement = statement.where(Expense.date < datetime.combine(date_to + timedelta(days=1), time.min))

    return statement


def attach_expense_allocations(session, rows):
    expense_ids = [row.id for row in rows]
    allocations_by_expense = {expense_id: [] for expense_id in expense_ids}
    if expense_ids:
        allocation_rows = session.execute(
            select(
                ExpenseAllocation,
                ExpenseBudgetCategory.name,
                ExpenseBudgetCategory.sort_order,
            )
            .join(
                ExpenseBudgetCategory,
                ExpenseBudgetCategory.id == ExpenseAllocation.category_id,
            )
            .where(ExpenseAllocation.expense_id.in_(expense_ids))
            .order_by(
                ExpenseAllocation.expense_id,
                ExpenseBudgetCategory.sort_order,
                ExpenseBudgetCategory.name,
            )
        ).all()
        for allocation, category_name, _sort_order in allocation_rows:
            allocations_by_expense.setdefault(allocation.expense_id, []).append(
                {
                    "category_id": allocation.category_id,
                    "category_name": category_name,
                    "amount": money(allocation.amount),
                    "source": allocation.source,
                    "details": allocation.details,
                }
            )

    for row in rows:
        row.allocations = allocations_by_expense.get(row.id, [])
        row.allocated_amount = sum(
            (allocation["amount"] for allocation in row.allocations),
            Decimal("0.00"),
        )
        row.unallocated_amount = money(row.amount) - row.allocated_amount
        row.receipt_items = expense_receipt_items(row)
        row.receipt_items_count = len(row.receipt_items)
    return rows


def build_category_totals(rows):
    totals = {}
    for row in rows:
        for allocation in row.allocations:
            category_id = allocation["category_id"]
            summary = totals.setdefault(
                category_id,
                {
                    "category_id": category_id,
                    "category_name": allocation["category_name"],
                    "amount": Decimal("0.00"),
                    "operations": 0,
                },
            )
            summary["amount"] += allocation["amount"]
            summary["operations"] += 1
    return sorted(totals.values(), key=lambda item: (-item["amount"], item["category_name"]))


def load_expense_report(filters):
    Session = make_session_factory()
    with Session() as session:
        categories = session.scalars(
            select(ExpenseBudgetCategory)
            .where(ExpenseBudgetCategory.active.is_(True))
            .order_by(
                ExpenseBudgetCategory.sort_order,
                ExpenseBudgetCategory.name,
            )
        ).all()
        if filters.get("category") == "unallocated":
            filters["category_label"] = "Не распределено"
        else:
            filters["category_label"] = next(
                (
                    category.name
                    for category in categories
                    if category.id == filters.get("category")
                ),
                filters.get("category", ""),
            )

        rows = session.scalars(
            build_expense_statement(filters).order_by(
                Expense.date.desc(), Expense.number.desc(), Expense.id
            )
        ).all()
        attach_expense_allocations(session, rows)
        as_of = session.scalar(select(func.max(Expense.synced_at)))

    allocated_total = sum(
        (row.allocated_amount for row in rows),
        Decimal("0.00"),
    )
    unallocated_total = sum(
        (max(row.unallocated_amount, Decimal("0.00")) for row in rows),
        Decimal("0.00"),
    )
    stats = {
        "count": len(rows),
        "total": sum((Decimal(row.amount or 0) for row in rows), Decimal("0")),
        "allocated_total": allocated_total,
        "unallocated_total": unallocated_total,
        "unallocated_count": sum(1 for row in rows if row.unallocated_amount > 0),
        "categories": len(
            {
                allocation["category_id"]
                for row in rows
                for allocation in row.allocations
            }
        ),
        "counterparties": len({row.counterparty for row in rows if row.counterparty}),
        "category_totals": build_category_totals(rows),
    }
    return rows, stats, categories, as_of


def load_expense_detail(expense_id):
    Session = make_session_factory()
    with Session() as session:
        expense = session.get(Expense, expense_id)
        if expense is None:
            return None, []
        attach_expense_allocations(session, [expense])
        categories = session.scalars(
            select(ExpenseBudgetCategory)
            .where(ExpenseBudgetCategory.active.is_(True))
            .order_by(
                ExpenseBudgetCategory.sort_order,
                ExpenseBudgetCategory.name,
            )
        ).all()
    expense.allocation_amounts = {
        allocation["category_id"]: allocation["amount"]
        for allocation in expense.allocations
    }
    return expense, categories


def build_monthly_expense_report(expense_rows, as_of):
    as_of_date = as_of.date() if isinstance(as_of, datetime) else as_of
    report_end = as_of_date.replace(day=1) - timedelta(days=1)
    months = {}

    for operation_date, amount, category, counterparty in expense_rows:
        if operation_date is None:
            continue
        operation_date = (
            operation_date.date()
            if isinstance(operation_date, datetime)
            else operation_date
        )
        if operation_date > report_end:
            continue

        key = operation_date.strftime("%Y-%m")
        row = months.setdefault(
            key,
            {
                "month": key,
                "period_end": None,
                "amount": Decimal("0"),
                "operations": 0,
                "categories": set(),
                "counterparties": set(),
                "change": None,
            },
        )
        row["amount"] += Decimal(amount or 0)
        row["operations"] += 1
        if category:
            row["categories"].add(category)
        if counterparty:
            row["counterparties"].add(counterparty)

    if not months:
        return [], {
            "months": 0,
            "total": Decimal("0"),
            "average": Decimal("0"),
            "operations": 0,
        }

    first_month = min(date.fromisoformat(f"{key}-01") for key in months)
    current = first_month
    while current <= report_end:
        key = current.strftime("%Y-%m")
        next_month = (current.replace(day=28) + timedelta(days=4)).replace(day=1)
        months.setdefault(
            key,
            {
                "month": key,
                "period_end": None,
                "amount": Decimal("0"),
                "operations": 0,
                "categories": set(),
                "counterparties": set(),
                "change": None,
            },
        )["period_end"] = next_month - timedelta(days=1)
        current = next_month

    rows = [months[key] for key in sorted(months)]
    previous_amount = None
    for row in rows:
        row["category_count"] = len(row.pop("categories"))
        row["counterparty_count"] = len(row.pop("counterparties"))
        if previous_amount is not None:
            row["change"] = row["amount"] - previous_amount
        previous_amount = row["amount"]

    total = sum((row["amount"] for row in rows), Decimal("0"))
    stats = {
        "months": len(rows),
        "total": total,
        "average": total / Decimal(len(rows)),
        "operations": sum(row["operations"] for row in rows),
    }
    rows.reverse()
    return rows, stats


def load_monthly_expense_report():
    Session = make_session_factory()
    with Session() as session:
        expense_rows = session.execute(
            select(
                Expense.date,
                Expense.amount,
                Expense.expense_category,
                Expense.counterparty,
            )
        ).all()
        as_of = session.scalar(select(func.max(Expense.synced_at))) or datetime.now()

    rows, stats = build_monthly_expense_report(expense_rows, as_of)
    latest = rows[0] if rows else None
    return rows, latest, stats, as_of


def allocation_label(row):
    parts = [
        f'{allocation["category_name"]}: {allocation["amount"]:.2f}'
        for allocation in getattr(row, "allocations", [])
    ]
    unallocated = getattr(row, "unallocated_amount", money(row.amount))
    if unallocated > 0:
        parts.append(f"Не распределено: {unallocated:.2f}")
    return "\n".join(parts) or "Не распределено"


def render_expenses_xlsx(rows, stats, filters, as_of):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Расходы"
    sheet.sheet_view.showGridLines = False

    dark_green = "176B5B"
    pale_green = "E8F3F0"
    border_gray = "D0D5DD"
    text_color = "1D2430"
    white = "FFFFFF"
    thin_gray = Side(style="thin", color=border_gray)

    sheet.merge_cells("A1:K1")
    sheet["A1"] = "Расходы ТСН «МИКРОРАЙОН ЭКОПАРК»"
    sheet["A1"].font = Font(size=16, bold=True, color=white)
    sheet["A1"].fill = PatternFill("solid", fgColor=dark_green)
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 30

    sheet["A3"] = "Данные обновлены"
    sheet["B3"] = as_of
    if as_of:
        sheet["B3"].number_format = "dd/mm/yyyy hh:mm"
    sheet["D3"] = "Строк"
    sheet["E3"] = stats["count"]
    sheet["G3"] = "Сумма"
    sheet["H3"] = float(stats["total"])
    sheet["H3"].number_format = '#,##0.00 [$₽-ru-RU]'

    filter_parts = []
    if filters.get("date_from"):
        filter_parts.append(f"с {filters['date_from']}")
    if filters.get("date_to"):
        filter_parts.append(f"по {filters['date_to']}")
    if filters.get("category"):
        filter_parts.append(
            f"статья: {filters.get('category_label') or filters['category']}"
        )
    if filters.get("q"):
        filter_parts.append(f"поиск: {filters['q']}")
    sheet.merge_cells("A5:K5")
    sheet["A5"] = "Фильтры: " + (", ".join(filter_parts) if filter_parts else "все расходы")
    sheet["A5"].font = Font(italic=True, color="667085")

    header_row = 7
    headers = [
        "№",
        "Дата",
        "Документ",
        "Номер",
        "Статьи сметы",
        "Статья 1С",
        "Контрагент",
        "Назначение",
        "Сумма",
        "Валюта",
        "Организация",
    ]
    for column, value in enumerate(headers, start=1):
        cell = sheet.cell(header_row, column, value)
        cell.font = Font(bold=True, color=white)
        cell.fill = PatternFill("solid", fgColor=dark_green)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(right=thin_gray, bottom=thin_gray)

    for index, row in enumerate(rows, start=1):
        values = [
            index,
            row.date,
            row.document,
            row.number,
            allocation_label(row),
            row.expense_category,
            row.counterparty,
            row.purpose,
            float(row.amount or 0),
            row.currency,
            row.organization,
        ]
        excel_row = header_row + index
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(excel_row, column, value)
            cell.border = Border(right=thin_gray, bottom=thin_gray)
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=column in {3, 5, 6, 7, 8, 11},
            )
            if excel_row % 2 == 1:
                cell.fill = PatternFill("solid", fgColor=pale_green)
        sheet.cell(excel_row, 2).number_format = "dd/mm/yyyy"
        sheet.cell(excel_row, 9).number_format = '#,##0.00 [$₽-ru-RU]'
        sheet.row_dimensions[excel_row].height = max(
            34,
            15 * max(1, allocation_label(row).count("\n") + 1),
        )

    table_end = max(header_row + len(rows), header_row + 1)
    if not rows:
        for column in range(1, len(headers) + 1):
            sheet.cell(header_row + 1, column, "")
    table = Table(displayName="ExpensesTable", ref=f"A{header_row}:K{table_end}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium4",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=False,
        showColumnStripes=False,
    )
    sheet.add_table(table)
    sheet.freeze_panes = f"F{header_row + 1}"
    sheet.auto_filter.ref = f"A{header_row}:K{table_end}"
    sheet.column_dimensions["A"].width = 8
    sheet.column_dimensions["B"].width = 14
    sheet.column_dimensions["C"].width = 24
    sheet.column_dimensions["D"].width = 14
    sheet.column_dimensions["E"].width = 38
    sheet.column_dimensions["F"].width = 24
    sheet.column_dimensions["G"].width = 28
    sheet.column_dimensions["H"].width = 42
    sheet.column_dimensions["I"].width = 16
    sheet.column_dimensions["J"].width = 10
    sheet.column_dimensions["K"].width = 30
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = f"{header_row}:{header_row}"

    for cell in sheet[3]:
        cell.font = Font(bold=True if cell.column in {1, 4, 7} else False, color=text_color)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
