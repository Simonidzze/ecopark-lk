import calendar
import re
from collections import defaultdict
from datetime import datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from sqlalchemy import func, select

from .db import make_session_factory
from .models import Accrual, Balance, OwnerPlot, Payment


XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

STATUS_LABELS = {
    "stopped": "Перестал платить",
    "never_paid": "Никогда не платил",
    "recent": "Платил в последний месяц",
    "under_month": "Долг до месяца",
    "no_accruals": "Нет начислений для оценки",
    "no_debt": "Без долга",
}


def decimal_value(value):
    return Decimal(value or 0)


def split_phones(value):
    phones = []
    for part in re.split(r"[,;\n\r]+", value or ""):
        phone = part.strip()
        if phone:
            phones.append(phone)
    return phones


def plot_sort_key(row):
    value = str(row.get("plot_number") or "").strip().lower().replace("ё", "е")
    match = re.match(r"^(\d+)\s*(.*)$", value)
    if not match:
        return (1, value)
    return (0, int(match.group(1)), match.group(2).strip())


def previous_calendar_month(value):
    year, month = (value.year - 1, 12) if value.month == 1 else (value.year, value.month - 1)
    day = min(value.day, calendar.monthrange(year, month)[1])
    return value.replace(year=year, month=month, day=day)


def payment_recency_band(last_payment_date, as_of):
    if last_payment_date is None:
        return "Нет оплат"
    if last_payment_date >= previous_calendar_month(as_of):
        return "До 1 месяца"
    days = (as_of - last_payment_date).days
    if days < 92:
        return "1–3 месяца"
    if days < 183:
        return "3–6 месяцев"
    if days < 365:
        return "6–12 месяцев"
    return "Более 12 месяцев"


def build_payment_status_report(owner_plot_rows, accrual_rows, payment_rows, as_of, scope="all"):
    cutoff = previous_calendar_month(as_of)
    owner_plot_to_owner = {
        row["owner_plot_id"]: row["owner_id"]
        for row in owner_plot_rows
    }

    owners = {}
    for row in owner_plot_rows:
        owner = owners.setdefault(
            row["owner_id"],
            {
                "owner_id": row["owner_id"],
                "owner": row["owner"],
                "phones": [],
                "plots": [],
                "overpayment": Decimal("0"),
                "total_debt": Decimal("0"),
            },
        )
        for phone in split_phones(row.get("phone")):
            if phone not in owner["phones"]:
                owner["phones"].append(phone)
        owner["plots"].append(
            {
                "plot_number": row.get("plot_number") or "",
                "account": row.get("account") or "",
            }
        )
        total = decimal_value(row.get("total"))
        if total > 0:
            owner["total_debt"] += total
        owner["overpayment"] += decimal_value(row.get("overpayment"))

    accruals_by_owner_month = defaultdict(lambda: defaultdict(Decimal))
    for owner_plot_id, operation_date, amount in accrual_rows:
        owner_id = owner_plot_to_owner.get(owner_plot_id)
        if not owner_id or operation_date is None:
            continue
        accruals_by_owner_month[owner_id][operation_date.strftime("%Y-%m")] += decimal_value(amount)

    payments_by_owner = defaultdict(list)
    for owner_plot_id, operation_date, amount in payment_rows:
        owner_id = owner_plot_to_owner.get(owner_plot_id)
        amount = decimal_value(amount)
        if not owner_id or operation_date is None or amount <= 0:
            continue
        payments_by_owner[owner_id].append((operation_date, amount))

    rows = []
    for owner_id, owner in owners.items():
        monthly_values = [
            amount
            for amount in accruals_by_owner_month.get(owner_id, {}).values()
            if amount > 0
        ]
        monthly_accrual = (
            sum(monthly_values, Decimal("0")) / Decimal(len(monthly_values))
            if monthly_values
            else None
        )
        debt_months = (
            owner["total_debt"] / monthly_accrual
            if monthly_accrual and monthly_accrual > 0
            else None
        )

        payments = sorted(payments_by_owner.get(owner_id, []), key=lambda item: item[0])
        first_payment_date = payments[0][0] if payments else None
        last_payment_date = payments[-1][0] if payments else None
        last_payment_amount = payments[-1][1] if payments else Decimal("0")
        total_paid = sum((amount for _date, amount in payments), Decimal("0"))
        payment_months = len({date.strftime("%Y-%m") for date, _amount in payments})

        if owner["total_debt"] <= 0:
            status_code = "no_debt"
        elif not payments:
            status_code = "never_paid"
        elif debt_months is None:
            status_code = "no_accruals"
        elif debt_months <= 1:
            status_code = "under_month"
        elif last_payment_date < cutoff:
            status_code = "stopped"
        else:
            status_code = "recent"

        plots = sorted(owner["plots"], key=plot_sort_key)
        rows.append(
            {
                "owner_id": owner_id,
                "owner": owner["owner"],
                "phones": owner["phones"],
                "plots": plots,
                "total_debt": owner["total_debt"],
                "overpayment": owner["overpayment"],
                "monthly_accrual": monthly_accrual,
                "debt_months": debt_months,
                "status_code": status_code,
                "status": STATUS_LABELS[status_code],
                "first_payment_date": first_payment_date,
                "last_payment_date": last_payment_date,
                "days_since_payment": (
                    (as_of - last_payment_date).days
                    if last_payment_date is not None
                    else None
                ),
                "payment_recency_band": payment_recency_band(last_payment_date, as_of),
                "last_payment_amount": last_payment_amount,
                "payment_count": len(payments),
                "payment_months": payment_months,
                "total_paid": total_paid,
                "average_payment": (
                    total_paid / Decimal(len(payments))
                    if payments
                    else None
                ),
            }
        )

    if scope == "stopped":
        rows = [row for row in rows if row["status_code"] == "stopped"]
    elif scope != "all":
        raise ValueError(f"Неизвестный состав отчета: {scope}")

    rows.sort(key=lambda row: (-row["total_debt"], row["owner"]))
    stats = {
        "count": len(rows),
        "total_debt": sum((row["total_debt"] for row in rows), Decimal("0")),
        "stopped": sum(row["status_code"] == "stopped" for row in rows),
        "never_paid": sum(row["status_code"] == "never_paid" for row in rows),
        "no_debt": sum(row["status_code"] == "no_debt" for row in rows),
    }
    return rows, stats, cutoff


def load_payment_status_report(scope="all"):
    Session = make_session_factory()
    with Session() as session:
        owner_plot_rows = session.execute(
            select(
                OwnerPlot.id.label("owner_plot_id"),
                OwnerPlot.owner_id,
                OwnerPlot.owner,
                OwnerPlot.phone,
                OwnerPlot.plot_number,
                OwnerPlot.account,
                Balance.total,
                Balance.overpayment,
            ).outerjoin(Balance, Balance.owner_plot_id == OwnerPlot.id)
        ).mappings().all()
        accrual_rows = session.execute(
            select(Accrual.owner_plot_id, Accrual.date, Accrual.amount)
        ).all()
        payment_rows = session.execute(
            select(Payment.owner_plot_id, Payment.date, Payment.amount)
            .where(Payment.date.is_not(None))
            .order_by(Payment.date)
        ).all()
        as_of = session.scalar(select(func.max(Balance.synced_at))) or datetime.now()

    rows, stats, cutoff = build_payment_status_report(
        owner_plot_rows,
        accrual_rows,
        payment_rows,
        as_of,
        scope=scope,
    )
    return rows, stats, as_of, cutoff


def render_payment_status_xlsx(rows, stats, as_of, cutoff, scope="all"):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Платежи владельцев"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "C13"

    dark_blue = "1F4E78"
    pale_blue = "D9EAF7"
    border_gray = "D0D5DD"
    text_color = "172B4D"
    white = "FFFFFF"
    thin_gray = Side(style="thin", color=border_gray)

    title = (
        "Владельцы, которые перестали платить"
        if scope == "stopped"
        else "Платежи всех владельцев"
    )
    sheet["A2"] = title
    sheet["A2"].font = Font(name="Arial", size=14, bold=True, color=dark_blue)
    for column in range(1, 20):
        sheet.cell(row=3, column=column).border = Border(
            bottom=Side(style="thin", color="5B9BD5")
        )

    sheet["A4"] = "Срез на"
    sheet["B4"] = as_of
    sheet["B4"].number_format = "dd/mm/yyyy hh:mm"
    sheet["D4"] = "Порог оплаты"
    sheet["E4"] = cutoff
    sheet["E4"].number_format = "dd/mm/yyyy hh:mm"
    sheet["A5"] = "Правило"
    sheet["B5"] = (
        "Статус «Перестал платить»: был положительный платеж, последний платеж старше "
        "одного календарного месяца, долг больше среднего месячного начисления."
    )
    sheet["B5"].font = Font(name="Arial", size=9, italic=True, color="667085")

    metric_headers = [
        "Владельцев",
        "Общий долг, ₽",
        "Перестали платить",
        "Никогда не платили",
        "Без долга",
    ]
    metric_values = [
        stats["count"],
        float(stats["total_debt"]),
        stats["stopped"],
        stats["never_paid"],
        stats["no_debt"],
    ]
    for column, value in enumerate(metric_headers, start=1):
        cell = sheet.cell(row=7, column=column, value=value)
        cell.fill = PatternFill("solid", fgColor=pale_blue)
        cell.font = Font(name="Arial", size=10, bold=True, color=dark_blue)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin_gray)
    for column, value in enumerate(metric_values, start=1):
        cell = sheet.cell(row=8, column=column, value=value)
        cell.font = Font(name="Arial", size=12, bold=True, color=text_color)
        cell.alignment = Alignment(horizontal="right", vertical="center")
        cell.border = Border(bottom=thin_gray)
    sheet["B8"].number_format = '#,##0.00;[Red](#,##0.00);-'

    sheet["A10"] = "Список для работы"
    sheet["A10"].font = Font(name="Arial", size=12, bold=True, color=dark_blue)
    sheet["A11"] = "Сортировка: по текущему долгу, от большего к меньшему. Источник: текущая база Ecopark."
    sheet["A11"].font = Font(name="Arial", size=9, italic=True, color="667085")

    headers = [
        "№",
        "Статус",
        "Владелец",
        "Телефон",
        "Участки",
        "Лицевые счета",
        "Текущий долг, ₽",
        "Переплата, ₽",
        "Среднее начисление в месяц, ₽",
        "Долг, мес.",
        "Последняя оплата",
        "Дней без оплаты",
        "Последняя сумма, ₽",
        "Первая оплата",
        "Оплат, шт.",
        "Месяцев с оплатами",
        "Всего оплачено, ₽",
        "Средний платеж, ₽",
        "Период без оплаты",
    ]
    header_row = 12
    for column, value in enumerate(headers, start=1):
        cell = sheet.cell(row=header_row, column=column, value=value)
        cell.fill = PatternFill("solid", fgColor=dark_blue)
        cell.font = Font(name="Arial", size=10, bold=True, color=white)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(right=Side(style="thin", color=white))

    status_fills = {
        "stopped": ("FDE8E7", "B42318"),
        "never_paid": ("FEF0C7", "93370D"),
        "recent": ("EEF4FF", "1849A9"),
        "under_month": ("F2F4F7", "475467"),
        "no_accruals": ("F2F4F7", "475467"),
        "no_debt": ("EDFCF2", "087443"),
    }
    first_data_row = header_row + 1
    for index, row in enumerate(rows, start=1):
        plots = ", ".join(plot["plot_number"] for plot in row["plots"] if plot["plot_number"])
        accounts = ", ".join(plot["account"] for plot in row["plots"] if plot["account"])
        values = [
            index,
            row["status"],
            row["owner"],
            ", ".join(row["phones"]),
            plots,
            accounts,
            float(row["total_debt"]),
            float(row["overpayment"]),
            float(row["monthly_accrual"]) if row["monthly_accrual"] is not None else None,
            float(row["debt_months"]) if row["debt_months"] is not None else None,
            row["last_payment_date"],
            row["days_since_payment"],
            float(row["last_payment_amount"]),
            row["first_payment_date"],
            row["payment_count"],
            row["payment_months"],
            float(row["total_paid"]),
            float(row["average_payment"]) if row["average_payment"] is not None else None,
            row["payment_recency_band"],
        ]
        excel_row = header_row + index
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=excel_row, column=column, value=value)
            cell.font = Font(name="Arial", size=10, color=text_color)
            cell.alignment = Alignment(vertical="center")
        fill_color, font_color = status_fills[row["status_code"]]
        sheet.cell(row=excel_row, column=2).fill = PatternFill("solid", fgColor=fill_color)
        sheet.cell(row=excel_row, column=2).font = Font(name="Arial", size=10, bold=True, color=font_color)

    last_data_row = header_row + len(rows)
    if rows:
        table = Table(displayName="PaymentStatusTable", ref=f"A{header_row}:S{last_data_row}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)

        sheet.conditional_formatting.add(
            f"L{first_data_row}:L{last_data_row}",
            CellIsRule(
                operator="greaterThanOrEqual",
                formula=["183"],
                fill=PatternFill("solid", fgColor="FDE8E7"),
                font=Font(color="B42318", bold=True),
            ),
        )
        sheet.conditional_formatting.add(
            f"L{first_data_row}:L{last_data_row}",
            CellIsRule(
                operator="between",
                formula=["92", "182"],
                fill=PatternFill("solid", fgColor="FEF0C7"),
                font=Font(color="93370D"),
            ),
        )

    currency_columns = (7, 8, 9, 13, 17, 18)
    for row_number in range(first_data_row, last_data_row + 1):
        sheet.cell(row=row_number, column=1).alignment = Alignment(horizontal="center", vertical="center")
        for column in currency_columns:
            sheet.cell(row=row_number, column=column).number_format = '#,##0.00;[Red](#,##0.00);-'
            sheet.cell(row=row_number, column=column).alignment = Alignment(horizontal="right", vertical="center")
        sheet.cell(row=row_number, column=10).number_format = "0.00"
        sheet.cell(row=row_number, column=11).number_format = "dd/mm/yyyy"
        sheet.cell(row=row_number, column=14).number_format = "dd/mm/yyyy"
        for column in (12, 15, 16):
            sheet.cell(row=row_number, column=column).number_format = "#,##0"
            sheet.cell(row=row_number, column=column).alignment = Alignment(horizontal="right", vertical="center")

    widths = [6, 24, 32, 23, 15, 19, 16, 16, 21, 12, 15, 16, 17, 15, 11, 18, 17, 17, 19]
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + column)].width = width
    sheet.row_dimensions[7].height = 30
    sheet.row_dimensions[12].height = 42
    for row_number in range(first_data_row, last_data_row + 1):
        sheet.row_dimensions[row_number].height = 22

    sheet.auto_filter.ref = f"A{header_row}:S{last_data_row}" if rows else f"A{header_row}:S{header_row}"
    sheet.print_options.horizontalCentered = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = f"1:{header_row}"

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
