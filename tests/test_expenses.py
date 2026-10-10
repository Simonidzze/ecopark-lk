import unittest
from datetime import datetime
from decimal import Decimal
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import load_workbook

from ecopark_sync.expense_report import (
    expense_filter_values,
    render_expenses_xlsx,
)
from ecopark_sync.payment_report import XLSX_CONTENT_TYPE
from ecopark_sync.models import Expense
from ecopark_sync.syncer import (
    linked_advance_report_items,
    linked_payroll_items,
    sync_expenses,
    synced_models_for_snapshot,
)
from ecopark_sync.web import create_app


class ExpenseSyncTest(unittest.TestCase):
    @patch("ecopark_sync.syncer.refresh_auto_expense_allocations_for_ids")
    @patch("ecopark_sync.syncer.upsert_many")
    def test_syncs_expense_fields(self, upsert_many, refresh_allocations):
        synced_at = datetime(2026, 9, 23, 10, 0)
        snapshot = {
            "expenses": [
                {
                    "id": "expense-1",
                    "document_id": "doc-1",
                    "document": "Списание с расчетного счета",
                    "date": "2026-09-22T12:00:00",
                    "number": "000123",
                    "expense_category_id": "category-1",
                    "expense_category": "Обслуживание территории",
                    "counterparty_id": "counterparty-1",
                    "counterparty": "ООО Подрядчик",
                    "purpose": "Покос травы",
                    "amount": 25000.50,
                    "currency": "RUB",
                    "organization_id": "org-1",
                    "organization": "ТСН Экопарк",
                    "source": "1C",
                }
            ]
        }

        session = object()
        sync_expenses(session, snapshot, 9, synced_at)

        rows = upsert_many.call_args.args[2]
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["id"], "expense-1")
        self.assertEqual(rows[0]["date"], datetime(2026, 9, 22, 12, 0))
        self.assertEqual(rows[0]["amount"], Decimal("25000.50"))
        self.assertEqual(rows[0]["expense_category"], "Обслуживание территории")
        self.assertEqual(rows[0]["receipt_items_json"], "[]")
        self.assertEqual(rows[0]["sync_run_id"], 9)
        self.assertEqual(
            upsert_many.call_args.kwargs["preserve_columns"],
            ("allocation_mode",),
        )
        refresh_allocations.assert_called_once_with(session, ["expense-1"])

    def test_links_only_unique_advance_report_item(self):
        snapshot = {
            "expenses": [
                {
                    "id": "expense-1",
                    "document_id": "doc-1",
                    "document": "Списание с расчетного счета 1",
                    "amount": "252.00",
                },
                {
                    "id": "expense-2",
                    "document_id": "doc-2",
                    "document": "Списание с расчетного счета 2",
                    "amount": "500.00",
                },
            ],
            "advance_report_advances": [
                {
                    "report_key": "Авансовый отчет 1",
                    "document": "Списание с расчетного счета 1",
                },
                {
                    "report_key": "Авансовый отчет 2",
                    "document": "Списание с расчетного счета 2",
                },
            ],
            "advance_report_items": [
                {
                    "report_key": "Авансовый отчет 1",
                    "report_number": "0001",
                    "receipt_line": 1,
                    "name": "Лента для ограждений",
                    "amount": 252,
                },
                {
                    "report_key": "Авансовый отчет 2",
                    "receipt_line": 1,
                    "name": "Кран",
                    "amount": 500,
                },
                {
                    "report_key": "Авансовый отчет 2",
                    "receipt_line": 2,
                    "name": "Труба",
                    "amount": 500,
                },
            ],
        }

        linked = linked_advance_report_items(snapshot)

        self.assertEqual(list(linked), ["doc-1"])
        self.assertEqual(linked["doc-1"][0]["name"], "Лента для ограждений")
        self.assertEqual(
            linked["doc-1"][0]["receipt_id"],
            "Авансовый отчет 1:1",
        )
        self.assertNotIn("report_key", linked["doc-1"][0])

    def test_links_collective_and_direct_payroll_expenses_without_duplication(self):
        snapshot = {
            "expenses": [
                {"document_id": "direct", "amount": "20.00"},
                {"document_id": "collective", "amount": "80.00"},
                {"document_id": "tax", "amount": "40.00"},
            ],
            "payroll_statement_items": [
                {
                    "statement_id": "statement-1",
                    "employee": "Ильиных Петр Сергеевич",
                    "amount": "20.00",
                    "period": "2026-09-01T00:00:00",
                },
                {
                    "statement_id": "statement-1",
                    "employee": "Дорогин Артем Владимирович",
                    "amount": "50.00",
                    "period": "2026-09-01T00:00:00",
                },
                {
                    "statement_id": "statement-1",
                    "employee": "Костенко Татьяна Гергартовна",
                    "amount": "30.00",
                    "period": "2026-09-01T00:00:00",
                },
            ],
            "payroll_expense_links": [
                {
                    "document_id": "direct",
                    "statement_id": "statement-1",
                    "employee": "Ильиных Петр Сергеевич",
                },
                {
                    "document_id": "collective",
                    "statement_id": "statement-1",
                    "employee": "",
                },
            ],
            "payroll_accrual_taxes": [
                {
                    "period": "2026-09-01T00:00:00",
                    "employee": "Дорогин Артем Владимирович",
                    "kind": "income_tax",
                    "amount": "30.00",
                },
                {
                    "period": "2026-09-01T00:00:00",
                    "employee": "Костенко Татьяна Гергартовна",
                    "kind": "insurance",
                    "amount": "10.00",
                },
            ],
            "payroll_tax_expenses": [
                {
                    "document_id": "tax",
                    "tax_period": "2026-09-01T00:00:00",
                    "tax": "Единый налоговый платеж",
                    "amount": "40.00",
                }
            ],
        }

        linked = linked_payroll_items(snapshot)

        self.assertEqual(len(linked["direct"]), 1)
        self.assertEqual(linked["direct"][0]["amount"], "20.00")
        self.assertEqual(
            {item["employee"] for item in linked["collective"]},
            {
                "Дорогин Артем Владимирович",
                "Костенко Татьяна Гергартовна",
            },
        )
        self.assertEqual(
            sum(Decimal(item["amount"]) for item in linked["collective"]),
            Decimal("80.00"),
        )
        self.assertEqual(
            sum(Decimal(item["amount"]) for item in linked["tax"]),
            Decimal("40.00"),
        )

    def test_old_snapshot_does_not_mark_expenses_as_stale(self):
        self.assertNotIn(Expense, synced_models_for_snapshot({}))
        self.assertIn(Expense, synced_models_for_snapshot({"expenses": []}))


class ExpenseReportTest(unittest.TestCase):
    def setUp(self):
        self.as_of = datetime(2026, 9, 23, 10, 0)
        self.rows = [
            SimpleNamespace(
                date=datetime(2026, 9, 22, 12, 0),
                document="Списание с расчетного счета",
                number="000123",
                expense_category="Обслуживание территории",
                counterparty="ООО Подрядчик",
                purpose="Покос травы",
                amount=Decimal("25000.50"),
                currency="RUB",
                organization="ТСН Экопарк",
            )
        ]
        self.stats = {
            "count": 1,
            "total": Decimal("25000.50"),
            "categories": 1,
            "counterparties": 1,
        }
        self.filters = expense_filter_values(
            {
                "date_from": "2026-09-01",
                "date_to": "2026-09-30",
                "category": "Обслуживание территории",
                "allocation_status": "unallocated",
                "q": "трава",
            }
        )

    def test_parses_filters_and_renders_typed_xlsx(self):
        self.assertEqual(self.filters["date_from_value"].isoformat(), "2026-09-01")
        self.assertEqual(self.filters["allocation_status"], "unallocated")
        report = render_expenses_xlsx(self.rows, self.stats, self.filters, self.as_of)
        workbook = load_workbook(report, data_only=False)
        sheet = workbook["Расходы"]

        self.assertEqual(sheet["A1"].value, "Расходы ТСН «МИКРОРАЙОН ЭКОПАРК»")
        self.assertEqual(sheet["B3"].value, self.as_of)
        self.assertEqual(sheet["A7"].value, "№")
        self.assertEqual(sheet["E7"].value, "Статьи сметы")
        self.assertEqual(sheet["F7"].value, "Статья 1С")
        self.assertEqual(sheet["B8"].value, datetime(2026, 9, 22, 12, 0))
        self.assertEqual(sheet["I8"].value, 25000.5)
        self.assertIn("ExpensesTable", sheet.tables)
        self.assertEqual(sheet.freeze_panes, "F8")


    def test_keeps_legacy_unallocated_category_link_working(self):
        filters = expense_filter_values({"category": "unallocated"})

        self.assertEqual(filters["category"], "")
        self.assertEqual(filters["allocation_status"], "unallocated")


class ExpenseRoutesTest(unittest.TestCase):
    def test_shows_page_and_downloads_current_filtered_xlsx(self):
        as_of = datetime(2026, 9, 23, 10, 0)
        stats = {
            "count": 0,
            "total": Decimal("0"),
            "allocated_total": Decimal("0"),
            "unallocated_total": Decimal("0"),
            "unallocated_count": 0,
            "categories": 0,
            "counterparties": 0,
            "category_totals": [],
        }
        with patch(
            "ecopark_sync.web.load_expense_report",
            return_value=(
                [],
                stats,
                [SimpleNamespace(id="electricity", name="Электроэнергия")],
                as_of,
            ),
        ) as loader:
            client = create_app().test_client()
            page = client.get("/admin/expenses?category=electricity&allocation_status=unallocated")
            response = client.get("/admin/expenses/report.xlsx?category=electricity&allocation_status=unallocated")

        self.assertEqual(page.status_code, 200)
        self.assertIn("Скачать Excel", page.get_data(as_text=True))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.mimetype, XLSX_CONTENT_TYPE)
        self.assertIn("expenses-2026-09-23.xlsx", response.headers["Content-Disposition"])
        self.assertGreater(len(response.data), 1000)
        self.assertEqual(loader.call_count, 2)
        self.assertEqual(loader.call_args.args[0]["category"], "electricity")
        self.assertEqual(
            loader.call_args.args[0]["allocation_status"],
            "unallocated",
        )


if __name__ == "__main__":
    unittest.main()
