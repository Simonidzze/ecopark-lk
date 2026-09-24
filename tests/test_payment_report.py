import unittest
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

from openpyxl import load_workbook

from ecopark_sync.payment_report import (
    XLSX_CONTENT_TYPE,
    build_payment_status_report,
    previous_calendar_month,
    render_payment_status_xlsx,
)
from ecopark_sync.web import create_app


class PaymentStatusReportTest(unittest.TestCase):
    def setUp(self):
        self.as_of = datetime(2026, 9, 23, 6, 0, 0)
        self.owner_plots = [
            self.owner_plot("stopped", "Иванов", "300"),
            self.owner_plot("never", "Петров", "200"),
            self.owner_plot("recent", "Сидоров", "250"),
            self.owner_plot("small", "Орлов", "80"),
            self.owner_plot("clear", "Смирнов", "0", overpayment="50"),
        ]
        self.accruals = [
            (f"plot-{owner_id}", datetime(2026, 7, 1), Decimal("100"))
            for owner_id in ("stopped", "never", "recent", "small", "clear")
        ]
        self.payments = [
            ("plot-stopped", datetime(2026, 6, 1), Decimal("100")),
            ("plot-recent", datetime(2026, 9, 10), Decimal("125")),
            ("plot-small", datetime(2026, 1, 10), Decimal("80")),
        ]

    @staticmethod
    def owner_plot(owner_id, owner, total, overpayment="0"):
        return {
            "owner_plot_id": f"plot-{owner_id}",
            "owner_id": owner_id,
            "owner": owner,
            "phone": f"+7 999 000-00-0{len(owner_id)}",
            "plot_number": str(len(owner_id)),
            "account": f"00{len(owner_id)}",
            "total": Decimal(total),
            "overpayment": Decimal(overpayment),
        }

    def test_previous_calendar_month_preserves_calendar_boundary(self):
        self.assertEqual(
            previous_calendar_month(datetime(2026, 3, 31, 6, 0)),
            datetime(2026, 2, 28, 6, 0),
        )

    def test_builds_all_payment_statuses_and_filters_stopped(self):
        rows, stats, cutoff = build_payment_status_report(
            self.owner_plots,
            self.accruals,
            self.payments,
            self.as_of,
        )

        rows_by_owner = {row["owner_id"]: row for row in rows}
        self.assertEqual(cutoff, datetime(2026, 8, 23, 6, 0))
        self.assertEqual(rows_by_owner["stopped"]["status_code"], "stopped")
        self.assertEqual(rows_by_owner["never"]["status_code"], "never_paid")
        self.assertEqual(rows_by_owner["recent"]["status_code"], "recent")
        self.assertEqual(rows_by_owner["small"]["status_code"], "under_month")
        self.assertEqual(rows_by_owner["clear"]["status_code"], "no_debt")
        self.assertEqual(stats["count"], 5)
        self.assertEqual(stats["total_debt"], Decimal("830"))
        self.assertEqual(stats["stopped"], 1)
        self.assertEqual(stats["never_paid"], 1)
        self.assertEqual(stats["no_debt"], 1)

        stopped_rows, stopped_stats, _cutoff = build_payment_status_report(
            self.owner_plots,
            self.accruals,
            self.payments,
            self.as_of,
            scope="stopped",
        )
        self.assertEqual([row["owner_id"] for row in stopped_rows], ["stopped"])
        self.assertEqual(stopped_stats["count"], 1)
        self.assertEqual(stopped_stats["total_debt"], Decimal("300"))

    def test_renders_xlsx_with_typed_dates_and_filterable_table(self):
        rows, stats, cutoff = build_payment_status_report(
            self.owner_plots,
            self.accruals,
            self.payments,
            self.as_of,
        )
        report = render_payment_status_xlsx(rows, stats, self.as_of, cutoff)
        workbook = load_workbook(report, data_only=False)
        sheet = workbook["Платежи владельцев"]

        self.assertEqual(sheet["A2"].value, "Платежи всех владельцев")
        self.assertEqual(sheet["B4"].value, self.as_of)
        self.assertEqual(sheet["A12"].value, "№")
        self.assertEqual(sheet["B12"].value, "Статус")
        self.assertEqual(sheet["S12"].value, "Период без оплаты")
        self.assertEqual(sheet.freeze_panes, "C13")
        self.assertIn("PaymentStatusTable", sheet.tables)
        self.assertEqual(sheet.max_row, 17)
        self.assertIsInstance(sheet["K13"].value, datetime)


class PaymentStatusDownloadRouteTest(unittest.TestCase):
    def test_downloads_fresh_xlsx_and_exposes_buttons(self):
        as_of = datetime(2026, 9, 23, 6, 0)
        cutoff = datetime(2026, 8, 23, 6, 0)
        stats = {
            "count": 0,
            "total_debt": Decimal("0"),
            "stopped": 0,
            "never_paid": 0,
            "no_debt": 0,
        }
        with patch(
            "ecopark_sync.web.load_payment_status_report",
            return_value=([], stats, as_of, cutoff),
        ) as loader, patch(
            "ecopark_sync.web.load_debtors_report",
            return_value=([], {"count": 0, "total_debt": Decimal("0")}, []),
        ):
            client = create_app().test_client()
            response = client.get("/admin/payments/report.xlsx?scope=all")
            page = client.get("/admin/debtors")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.mimetype, XLSX_CONTENT_TYPE)
        self.assertIn("payment-report-all-2026-09-23.xlsx", response.headers["Content-Disposition"])
        self.assertGreater(len(response.data), 1000)
        loader.assert_called_once_with(scope="all")
        page_text = page.get_data(as_text=True)
        self.assertIn("Скачать всех в Excel", page_text)
        self.assertIn("Перестали платить", page_text)

    def test_rejects_unknown_scope(self):
        client = create_app().test_client()
        response = client.get("/admin/payments/report.xlsx?scope=unknown")
        self.assertEqual(response.status_code, 400)


if __name__ == "__main__":
    unittest.main()
